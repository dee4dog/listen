"""The Word-template exporter.

The awkward parts are all in how Word actually stores a document: a
placeholder is routinely split across runs, fields appear in headers and
table cells as well as the body, and a block field has to expand without
losing the style of the paragraph it replaced.
"""
import pytest
from docx import Document

from scribe.exporters.docx_template import (KNOWN_FIELDS, export_docx_template,
                                            scan_template,
                                            write_starter_template)


@pytest.fixture
def starter(tmp_path):
    path = tmp_path / "starter.docx"
    write_starter_template(path)
    return path


def body_text(path):
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


# ------------------------------------------------------------ the starter

def test_starter_template_uses_only_known_fields(starter):
    """It is the user's first example, so it must not warn about itself."""
    found, unknown = scan_template(starter)
    assert not unknown
    assert {"title", "summary", "transcript", "action_items"} <= found
    assert found <= KNOWN_FIELDS


def test_starter_template_fills_completely(starter, session, tmp_path):
    out = tmp_path / "filled.docx"
    export_docx_template(session, starter, out)
    text = body_text(out)

    assert "{{" not in text, "a placeholder was left behind"
    assert "Site meeting — Block C" in text
    assert "25 August 2026" in text          # ISO date rendered for people
    assert "01:12:40" in text                # duration formatted
    assert "Sarah, Johan" in text
    assert "[00:00:05] Johan: Ek sal die verslag teen Vrydag stuur." in text
    assert "RECURRING — first raised 2026-07-14" in text


def test_export_does_not_modify_the_template(starter, session, tmp_path):
    before = starter.read_bytes()
    export_docx_template(session, starter, tmp_path / "out.docx")
    assert starter.read_bytes() == before


# ------------------------------------------------- how Word stores things

@pytest.fixture
def tricky(tmp_path):
    """A template with the shapes that break naive replacement."""
    doc = Document()

    para = doc.add_paragraph()
    para.add_run("Minutes for: ").bold = True   # formatting around a field
    para.add_run("{{ti")                        # Word splits placeholders
    para.add_run("tle}} on {{date}}")           # after a mid-word edit

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Duration"
    table.cell(0, 1).text = "{{duration}}"
    table.cell(1, 0).text = "Present"
    table.cell(1, 1).text = "{{participants}}"

    doc.sections[0].header.paragraphs[0].text = "{{app}} — {{discussion_type}}"
    doc.sections[0].footer.paragraphs[0].text = "Exported {{generated_on}}"

    doc.add_paragraph("{{decisions}}", style="List Bullet")
    doc.add_paragraph("{{nonsense_field}}")

    path = tmp_path / "tricky.docx"
    doc.save(str(path))
    return path


def test_placeholder_split_across_runs_is_still_replaced(tricky, session,
                                                         tmp_path):
    out = tmp_path / "out.docx"
    export_docx_template(session, tricky, out)
    first = Document(str(out)).paragraphs[0]
    assert first.text == "Minutes for: Site meeting — Block C on 25 August 2026"


def test_formatting_around_a_field_survives(tricky, session, tmp_path):
    out = tmp_path / "out.docx"
    export_docx_template(session, tricky, out)
    first = Document(str(out)).paragraphs[0]
    label = first.runs[0]
    assert label.text == "Minutes for: " and label.bold is True


def test_fields_in_tables_headers_and_footers(tricky, session, tmp_path):
    out = tmp_path / "out.docx"
    export_docx_template(session, tricky, out)
    doc = Document(str(out))

    assert doc.tables[0].cell(0, 1).text == "01:12:40"
    assert "Sarah, Johan" in doc.tables[0].cell(1, 1).text
    assert "Listen — Meeting" in doc.sections[0].header.paragraphs[0].text
    assert "Exported" in doc.sections[0].footer.paragraphs[0].text


def test_block_field_keeps_its_paragraph_style(tricky, session, tmp_path):
    """A bulleted {{decisions}} must produce bulleted decisions."""
    out = tmp_path / "out.docx"
    export_docx_template(session, tricky, out)
    bullets = [p for p in Document(str(out)).paragraphs
               if p.style.name == "List Bullet"]
    assert bullets and "Approved the revised layout" in bullets[0].text


def test_unknown_field_is_reported_and_left_alone(tricky, session, tmp_path):
    """Leaving a typo visible beats making it disappear silently."""
    _found, unknown = scan_template(tricky)
    assert unknown == {"nonsense_field"}

    out = tmp_path / "out.docx"
    export_docx_template(session, tricky, out)
    assert "{{nonsense_field}}" in body_text(out)


# --------------------------------------------------------- odds and ends

def test_spacing_and_case_inside_the_braces(tmp_path, session):
    doc = Document()
    doc.add_paragraph("{{ Title }} / {{TITLE}} / {{title}}")
    template = tmp_path / "case.docx"
    doc.save(str(template))

    out = tmp_path / "out.docx"
    export_docx_template(session, template, out)
    assert body_text(out) == " / ".join(["Site meeting — Block C"] * 3)


def test_empty_section_says_so_rather_than_leaving_a_blank(tmp_path, session):
    session.issues = [i for i in session.issues if i.kind != "decision"]
    doc = Document()
    doc.add_paragraph("{{decisions}}")
    template = tmp_path / "empty.docx"
    doc.save(str(template))

    out = tmp_path / "out.docx"
    export_docx_template(session, template, out)
    assert body_text(out).strip() == "None recorded."


def test_block_field_expands_to_one_paragraph_per_line(tmp_path, session):
    doc = Document()
    doc.add_paragraph("{{transcript_plain}}")
    template = tmp_path / "lines.docx"
    doc.save(str(template))

    out = tmp_path / "out.docx"
    export_docx_template(session, template, out)
    lines = [p.text for p in Document(str(out)).paragraphs if p.text]
    assert len(lines) == len(session.segments)
    assert lines[0] == "Sarah: Morning everyone, let's start."
    assert "[00:00:00]" not in lines[0]      # the plain variant drops stamps


def test_scan_rejects_a_file_that_is_not_a_document(tmp_path):
    bogus = tmp_path / "not.docx"
    bogus.write_bytes(b"this is not a zip archive")
    with pytest.raises(Exception):
        scan_template(bogus)
