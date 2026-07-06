"""Word (.docx) export via python-docx."""
from docx import Document
from docx.shared import Pt, RGBColor

from ..analysis.discussion import DISCUSSION_TYPES, SECTION_TITLES, items_for_section
from ..models import fmt_ts

_RED = RGBColor(0xB3, 0x00, 0x00)
_GREY = RGBColor(0x66, 0x66, 0x66)


def export_docx(session, path):
    dtype = DISCUSSION_TYPES.get(session.dtype, DISCUSSION_TYPES["general"])
    doc = Document()
    doc.add_heading(session.title, level=0)
    doc.add_paragraph(
        f"{dtype['label']}  |  {session.started_at.replace('T', ' ')}  |  "
        f"Duration: {fmt_ts(session.duration)}")
    doc.add_paragraph(f"Participants: {', '.join(session.speaker_names()) or '—'}")

    for section in (session.sections_override or dtype["sections"]):
        items = items_for_section(session.issues, section)
        if not items:
            continue
        doc.add_heading(SECTION_TITLES[section], level=1)
        for item in items:
            para = doc.add_paragraph(style="List Bullet")
            if item.speaker:
                para.add_run(f"{item.speaker}: ").bold = True
            para.add_run(item.text)
            if item.recurring:
                flag = para.add_run(
                    f"  [RECURRING — first raised {item.prior_date or 'earlier'}]")
                flag.bold = True
                flag.font.color.rgb = _RED

    doc.add_heading("Transcript", level=1)
    for seg in session.segments:
        para = doc.add_paragraph()
        ts_run = para.add_run(f"[{fmt_ts(seg.start)}] ")
        ts_run.font.color.rgb = _GREY
        ts_run.font.size = Pt(9)
        para.add_run(f"{seg.speaker}: ").bold = True
        para.add_run(seg.text)

    doc.save(str(path))
    return path
