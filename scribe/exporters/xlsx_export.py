"""Excel (.xlsx) export via openpyxl: Summary, Transcript and Speakers sheets."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..analysis.discussion import DISCUSSION_TYPES, SECTION_TITLES, items_for_section
from ..models import fmt_ts

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_RECURRING_FILL = PatternFill("solid", fgColor="FFC7CE")


def _header(ws, columns, widths):
    ws.append(columns)
    for i, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        ws.column_dimensions[get_column_letter(i)].width = width


def export_xlsx(session, path):
    dtype = DISCUSSION_TYPES.get(session.dtype, DISCUSSION_TYPES["general"])
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Title", session.title])
    ws.append(["Type", dtype["label"]])
    ws.append(["Date", session.started_at.replace("T", " ")])
    ws.append(["Duration", fmt_ts(session.duration)])
    ws.append(["Participants", ", ".join(session.speaker_names()) or "—"])
    for row in range(1, 6):
        ws.cell(row=row, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 90
    ws.append([])
    ws.append(["Section", "Speaker", "Item", "Recurring", "First raised"])
    header_row = ws.max_row
    for col in range(1, 6):
        cell = ws.cell(row=header_row, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for section in (session.sections_override or dtype["sections"]):
        for item in items_for_section(session.issues, section):
            ws.append([SECTION_TITLES[section], item.speaker, item.text,
                       "YES" if item.recurring else "",
                       item.prior_date or ""])
            if item.recurring:
                for col in range(1, 6):
                    ws.cell(row=ws.max_row, column=col).fill = _RECURRING_FILL
    ws.column_dimensions["C"].width = 90

    ws2 = wb.create_sheet("Transcript")
    _header(ws2, ["Start", "End", "Speaker", "Text"], [10, 10, 20, 120])
    for seg in session.segments:
        ws2.append([fmt_ts(seg.start), fmt_ts(seg.end), seg.speaker, seg.text])

    ws3 = wb.create_sheet("Speakers")
    _header(ws3, ["Speaker", "Segments", "Speaking time"], [24, 12, 16])
    stats = {}
    for seg in session.segments:
        count, secs = stats.get(seg.speaker, (0, 0.0))
        stats[seg.speaker] = (count + 1, secs + max(0.0, seg.end - seg.start))
    for name in session.speaker_names():
        count, secs = stats.get(name, (0, 0.0))
        ws3.append([name, count, fmt_ts(secs)])

    wb.save(str(path))
    return path
