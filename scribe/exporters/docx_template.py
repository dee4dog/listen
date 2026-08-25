"""Word export driven by a user-supplied .docx template.

The user designs an ordinary Word document — their letterhead, fonts, tables,
headers and footers — and drops **placeholders** such as ``{{title}}`` or
``{{transcript}}`` where the session's content should appear. Listen fills
them in and saves the result, so exports always look like the organisation's
own paperwork.

Two kinds of placeholder:

* **Inline fields** are replaced in place, anywhere in the document (body,
  tables, headers, footers). The surrounding formatting is preserved, and the
  inserted value takes on the formatting of the placeholder itself.
* **Block fields** stand alone on their own paragraph and expand into as many
  paragraphs as the content needs. Every generated paragraph inherits the
  placeholder paragraph's style, so styling ``{{action_items}}`` as a bullet
  list turns every action item into a bullet.

No template language, no loops, no conditionals — just the names below.
"""
import re
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from ..analysis.discussion import (DISCUSSION_TYPES, SECTION_TITLES,
                                   items_for_section)
from ..models import fmt_ts

# A placeholder: {{name}}, tolerant of spaces and capitals — {{ Title }} works.
_FIELD_RE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)\s*\}\}")

# Shown to the user in the "How to build a template" help and in the README.
INLINE_FIELDS = [
    ("title", "Session title, e.g. “Site meeting 12 March”"),
    ("discussion_type", "Meeting, Brief, General discussion or Interview"),
    ("date", "Date of the session, e.g. “25 August 2026”"),
    ("time", "Start time, e.g. “14:03”"),
    ("datetime", "Date and time together, e.g. “2026-08-25 14:03”"),
    ("duration", "Length of the recording, e.g. “01:12:40”"),
    ("participants", "All speaker names on one line, comma separated"),
    ("speaker_count", "How many speakers were detected"),
    ("line_count", "How many transcript lines there are"),
    ("word_count", "How many words were spoken in total"),
    ("generated_on", "Date and time this document was exported"),
    ("app", "The name of the app that produced the document"),
]

BLOCK_FIELDS = [
    ("summary", "The whole summary: every section for this discussion type, "
                "with its heading"),
    ("action_items", "One paragraph per action item / task"),
    ("decisions", "One paragraph per decision"),
    ("key_issues", "One paragraph per key issue (recurring ones are flagged)"),
    ("key_points", "One paragraph per key point"),
    ("participants_list", "One paragraph per speaker"),
    ("transcript", "The full transcript: “[00:01:23] Name: what they said”"),
    ("transcript_plain", "The transcript without timestamps: “Name: …”"),
]

_SECTION_FOR_BLOCK = {
    "action_items": "actions",
    "decisions": "decisions",
    "key_issues": "issues",
    "key_points": "key_points",
}

_OTHER_BLOCKS = {"summary", "transcript", "transcript_plain",
                 "participants_list"}

_EMPTY = "None recorded."

KNOWN_FIELDS = ({name for name, _ in INLINE_FIELDS}
                | {name for name, _ in BLOCK_FIELDS})


# --------------------------------------------------------------- the values

def _inline_values(session):
    from .. import APP_NAME
    try:
        started = datetime.fromisoformat(session.started_at)
    except (TypeError, ValueError):
        started = None
    words = sum(len(seg.text.split()) for seg in session.segments)
    dtype = DISCUSSION_TYPES.get(session.dtype, DISCUSSION_TYPES["general"])
    names = session.speaker_names()
    return {
        "title": session.title,
        "discussion_type": dtype["label"],
        "date": f"{started:%d %B %Y}" if started else session.started_at[:10],
        "time": f"{started:%H:%M}" if started else "",
        "datetime": session.started_at.replace("T", " "),
        "duration": fmt_ts(session.duration),
        "participants": ", ".join(names) or "—",
        "speaker_count": str(len(names)),
        "line_count": str(len(session.segments)),
        "word_count": str(words),
        "generated_on": f"{datetime.now():%d %B %Y %H:%M}",
        "app": APP_NAME,
    }


def _item_chunks(item):
    """One summary item as (text, bold) chunks."""
    chunks = []
    if item.speaker:
        chunks.append((f"{item.speaker}: ", True))
    chunks.append((item.text, None))
    if item.recurring:
        chunks.append((f"  [RECURRING — first raised "
                       f"{item.prior_date or 'earlier'}]", True))
    return chunks


def _section_lines(session, section):
    items = items_for_section(session.issues, section)
    if not items:
        return [[(_EMPTY, None)]]
    return [_item_chunks(item) for item in items]


def _block_lines(session, name):
    """The paragraphs a block placeholder expands into, each a chunk list."""
    if name in _SECTION_FOR_BLOCK:
        return _section_lines(session, _SECTION_FOR_BLOCK[name])

    if name == "summary":
        dtype = DISCUSSION_TYPES.get(session.dtype, DISCUSSION_TYPES["general"])
        lines = []
        for section in (session.sections_override or dtype["sections"]):
            items = items_for_section(session.issues, section)
            if not items:
                continue
            lines.append([(SECTION_TITLES[section], True)])
            lines.extend(_item_chunks(item) for item in items)
        return lines or [[(_EMPTY, None)]]

    if name == "participants_list":
        names = session.speaker_names()
        return [[(n, None)] for n in names] or [[("—", None)]]

    if name in ("transcript", "transcript_plain"):
        stamps = name == "transcript"
        lines = []
        for seg in session.segments:
            chunks = []
            if stamps:
                chunks.append((f"[{fmt_ts(seg.start)}] ", None))
            chunks.append((f"{seg.speaker}: ", True))
            chunks.append((seg.text, None))
            lines.append(chunks)
        return lines or [[("No transcript.", None)]]

    return []


# ------------------------------------------------------------ docx plumbing

def _iter_paragraphs(container):
    """Every paragraph in a document part, including inside (nested) tables."""
    for para in container.paragraphs:
        yield para
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _all_paragraphs(doc):
    yield from _iter_paragraphs(doc)
    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part is not None:
                yield from _iter_paragraphs(part)


def _set_chunks(para, chunks):
    """Replace a paragraph's runs with `chunks`, keeping its character
    formatting: every new run is a clone of the paragraph's first run."""
    runs = para.runs
    donor = deepcopy(runs[0]._r) if runs else None
    for run in runs:
        run._element.getparent().remove(run._element)
    for text, bold in chunks:
        if donor is None:
            run = para.add_run(text)
        else:
            element = deepcopy(donor)
            para._p.append(element)
            run = Run(element, para)
            run.text = text
        if bold is not None:
            run.bold = bold


def _replace_inline(para, values, used):
    """Substitute inline fields, even where Word has split a placeholder
    across several runs. Text outside the placeholders stays in the runs it
    came from, so surrounding bold/italic survives."""
    runs = para.runs
    if not runs:
        return
    text = "".join(run.text for run in runs)
    matches = [m for m in _FIELD_RE.finditer(text)
               if m.group(1).lower() in values]
    if not matches:
        return

    spans, pos = [], 0
    for run in runs:
        spans.append((pos, pos + len(run.text)))
        pos += len(run.text)

    def run_at(offset):
        for index, (start, end) in enumerate(spans):
            if start <= offset < end:
                return index
        return len(runs) - 1

    parts = [""] * len(runs)

    def add_literal(start, end):
        for index, (run_start, run_end) in enumerate(spans):
            lo, hi = max(start, run_start), min(end, run_end)
            if lo < hi:
                parts[index] += text[lo:hi]

    cursor = 0
    for match in matches:
        add_literal(cursor, match.start())
        name = match.group(1).lower()
        parts[run_at(match.start())] += values[name]
        used.add(name)
        cursor = match.end()
    add_literal(cursor, len(text))

    for run, new_text in zip(runs, parts):
        run.text = new_text


def _expand_block(para, lines):
    """Replace a lone block placeholder with one paragraph per line, each a
    copy of the placeholder paragraph so its style and numbering carry over."""
    anchor = para._p
    previous = anchor
    for chunks in lines:
        element = deepcopy(anchor)
        previous.addnext(element)
        previous = element
        _set_chunks(Paragraph(element, para._parent), chunks)
    anchor.getparent().remove(anchor)


def _block_name(para):
    """The block field, if this paragraph holds nothing but one placeholder."""
    match = _FIELD_RE.fullmatch(para.text.strip())
    if match is None:
        return None
    name = match.group(1).lower()
    if name in _SECTION_FOR_BLOCK or name in _OTHER_BLOCKS:
        return name
    return None


# ------------------------------------------------------------------- public

def scan_template(path):
    """Look over a template before it is used.

    Returns (found, unknown): the recognised placeholder names it contains and
    the ones Listen does not know about — usually a typo worth telling the
    user about.
    """
    doc = Document(str(path))
    found, unknown = set(), set()
    for para in _all_paragraphs(doc):
        for match in _FIELD_RE.finditer(para.text):
            name = match.group(1).lower()
            (found if name in KNOWN_FIELDS else unknown).add(name)
    return found, unknown


def export_docx_template(session, template_path, out_path):
    """Fill `template_path` with `session` and save it as `out_path`."""
    doc = Document(str(template_path))
    values = _inline_values(session)
    used = set()

    # Blocks first: they are whole paragraphs, and expanding them must not
    # disturb the inline pass.
    for para in list(_all_paragraphs(doc)):
        name = _block_name(para)
        if name is not None:
            _expand_block(para, _block_lines(session, name))
            used.add(name)

    for para in _all_paragraphs(doc):
        _replace_inline(para, values, used)

    doc.save(str(out_path))
    return out_path


def write_starter_template(path):
    """Write a ready-to-edit template that uses the common placeholders, so
    the user can open it in Word, restyle it, and save it as their own."""
    from .. import APP_NAME
    doc = Document()
    doc.add_heading("{{title}}", level=0)
    doc.add_paragraph(
        "{{discussion_type}}  |  {{date}} at {{time}}  |  "
        "Duration {{duration}}")
    doc.add_paragraph("Participants: {{participants}} "
                      "({{speaker_count}} speakers, {{word_count}} words)")

    doc.add_heading("Summary", level=1)
    doc.add_paragraph("{{summary}}")

    doc.add_heading("Action items", level=1)
    doc.add_paragraph("{{action_items}}", style="List Bullet")

    doc.add_heading("Key issues", level=1)
    doc.add_paragraph("{{key_issues}}", style="List Bullet")

    doc.add_heading("Transcript", level=1)
    doc.add_paragraph("{{transcript}}")

    # The note deliberately avoids writing a placeholder of its own, so that
    # checking this template never reports an unknown field.
    note = doc.add_paragraph()
    note.add_run(
        f"\nThis is a starter template for {APP_NAME}. Restyle it however you "
        f"like — change the fonts and colours, add your logo, move fields into "
        f"the header or into a table — and keep each field in double braces "
        f"wherever you want that content to appear. Delete any you do not "
        f"need, including this note."
    ).italic = True
    doc.save(str(path))
    return Path(path)
